import csv
import openpyxl

from openai import OpenAI
from ML import getscore




#To change Nomination of headers
def headers_processed(L):
    H=[]
    for i in L:
        if i=='tone':
            H.append('Tone')
        elif i=='Message Clarity':
            H.append('Message')
        elif i=='effective use of tagging':
            H.append('Hashtag_Effectiveness')
        elif i=='Use of Controversial Topics':
            H.append('Controversial_Topics')
        elif i=='How well Call to actions are implemented':
            H.append('Call_To_Action')
        elif i=='how well the post is engaging':
            H.append('Engagement')
        else:
            H.append(i)
    return(H)


#Save a list of lists to an Excel FIle
def save_to_excel(data, filename, headers):
    """
    Save a list of lists to an Excel file.

    Args:
    data (list of lists): The data to be saved.
    filename (str): The name of the Excel file.
    """
    # Create a new workbook and select the active worksheet
    workbook = openpyxl.Workbook()
    sheet = workbook.active

    # add headers to the sheet:
    sheet.append(headers)

    # Write data to the worksheet
    for row in data:
        row.append("Suggestions")
        sheet.append(row)
    print(row)
    # Save the workbook to the specified file
    workbook.save(filename)
#Convert a CSV to a List of lists where each list Item is a List element
def CSVtoList(csvpath):
    # List to store the rows
    rows_list = []

    # Read the CSV file
    with open(csvpath, mode='r', encoding='utf-8') as file:
        # Create a CSV reader object
        csv_reader = csv.reader(file)
        
        # Iterate over each row in the CSV file
        for row in csv_reader:
            # Append each row to the list
            rows_list.append(row)

    return(rows_list)
#Count the Number of Hashtags in a given string
def count_hashtags(text):
    """
    Count the number of hashtags in the given text.

    Args:
    text (str): The input text containing hashtags.

    Returns:
    int: The number of hashtags in the text.
    """
    # Split the text into words
    words = text.split()
    
    # Count the words that start with '#'
    hashtag_count = sum(1 for word in words if word.startswith('#'))
    
    return hashtag_count
#Count the length of a post in a given string
def calculate_post_length(post):
    """
    Calculate the length of a post.

    Args:
    post (str): The input text of the post.

    Returns:
    int: The length of the post in characters.
    """
    return len(post)
#Helper Functions
#Calculate the occurrences of each word in a list of words
def calculate_word_occurrences(word_list):
    """
    Calculate the occurrences of each word in a list of words.

    Args:
    word_list (list): The list of words.

    Returns:
    list of lists: A list where each sublist contains a word and its count.
    """
    # Create an empty dictionary to store word counts
    word_counts = {}
    
    # Iterate over each word in the list
    for word in word_list:
        # Increment the count for the word
        if word in word_counts:
            word_counts[word] += 1
        else:
            word_counts[word] = 1
    
    # Convert the dictionary to a list of lists
    occurrences_list = [[word, count] for word, count in word_counts.items()]
    
    return occurrences_list
#extract all osts in postsuggestions in a long string
def extract_text_frame(suggestionfilepath):
    #Create a List with content only 
    Texttrame=""
    for row in CSVtoList(suggestionfilepath):
        Texttrame+=row[7]+' '
    return(Texttrame)
#Add newer column to posstsuggestion.csv
def enrich_input(suggestionInputFilepath,suggestionOutputFilepath):
    #Declare top keywords
    top_keywords=filter_top_occurrences(calculate_word_occurrences(extract_text_frame(suggestionInputFilepath).split(' ')))
    # Print the list to verify
    final=[]
    lo=[]
    count =0
    for row in CSVtoList(suggestionInputFilepath):
        H=[]
        count = count + 1
        for i in range(len(row)):
            H.append(row[i])
            #Hashtag Count
        
        H.append(count_hashtags(row[7]))
        H.append(calculate_post_length(row[7]))
        lo=(common_words_with_count(calculate_word_occurrences(row[7].split(' ')),top_keywords))
        H.append(str((lo)))
        H.append(keyscore(lo))

        # saving the Ml metric of the first three rows to the data
        if count <=3:
            ml_header=[]
            ml_metric =save_ml_metric(row[i])
            if ml_metric != []:
                for t in ml_metric:
                    H.append((t[1]))
                    if t[0] not in ml_header:
                        ml_header.append(t[0])

        final.append(H)
    # Save the data to an Excel file
    

    headers = ["post_id", "likes"	,"comments","shares","previous_likes","like_change","post_type","caption","Hash_Tag_Counts", "Post_Length","Common_Words_Count", "Keyscore"]
    headers = headers + ml_header
    headers=headers_processed(headers)+["Suggestions"]
    save_to_excel(final, suggestionOutputFilepath,headers)
#Keep only most used keywords
def filter_top_occurrences(lst):
    # Sort the list based on the occurrence count
    sorted_lst = sorted(lst, key=lambda x: x[1], reverse=True)
    
    # Calculate the number of elements to keep (top 20%)
    n = len(sorted_lst)
    k = int(n * 0.2)
    
    # Return the top k elements
    return sorted_lst[:k]
def common_words_with_count(list1, list2):
    # Convert lists to dictionaries for easier comparison
    dict1 = dict(list1)
    dict2 = dict(list2)
    
    # Initialize result list
    result = []
    
    # Iterate over words in the first list
    for word, count in list1:
        # Check if the word exists in the second list
        if word in dict2:
            # Add the word and its count from the first list to the result
            result.append([word, count])
    
    return result
#A function to calculate a keyword score
def keyscore(lst):
    key_score=0
    for i in lst:
        key_score+=i[1]
    return(key_score+len(lst))
#Save ml_Metric
def save_ml_metric(caption):
    Allscore=getscore(caption)
    if Allscore != []:
        return Allscore
    else:
        return []
#Tests
enrich_input('before.csv','after.xlsx')

